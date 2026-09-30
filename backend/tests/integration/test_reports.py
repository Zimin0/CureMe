"""Выписка для врача в PDF и Excel: период, чья история, комментарии, часовой пояс, Плюс."""

from datetime import datetime, timedelta, timezone
from io import BytesIO
from urllib.parse import unquote

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.doctor_report import DISCLAIMER, NO_COMMENTS_NOTE
from app.models import Intake
from tests.conftest import fid, register


@pytest.fixture
def fam(client):
    """Семья: владелец (первый аккаунт — админ) и мама; оба что-то принимали."""
    h, u = register(client)
    f = fid(u)
    code = client.get(f"/api/families/{f}", headers=h).json()["invite_code"]
    h_mom, mom = register(client, "mom@example.com", "Мама", invite=code)
    med = client.post(f"/api/families/{f}/medicines", headers=h, json={
        "name": "Нурофен", "dosage": "200 мг", "active_ingredient": "ибупрофен", "unit": "таб",
        "packages": [{"quantity": 20}]}).json()
    client.post(f"/api/families/{f}/medicines/{med['id']}/consume", headers=h, json={"amount": 1, "comment": "болела голова"})
    client.post(f"/api/families/{f}/medicines/{med['id']}/consume", headers=h_mom, json={"amount": 0.5, "comment": "мамин секрет"})
    return {"h": h, "u": u, "f": f, "h_mom": h_mom, "mom": mom, "med": med}


def xlsx(client, h, f, **params):
    r = client.get(f"/api/families/{f}/report.xlsx", headers=h, params=params)
    assert r.status_code == 200, r.text
    return load_workbook(BytesIO(r.content))


def rows(ws):
    return [r for r in ws.iter_rows(min_row=5, values_only=True)]


def test_pdf_has_cyrillic_font_and_file_name(client, fam):
    r = client.get(f"/api/families/{fam['f']}/report.pdf", headers=fam["h"])
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF-")
    assert b"DejaVu" in r.content  # шрифт с кириллицей встроен, а не стандартный Helvetica без неё
    cd = unquote(r.headers["content-disposition"])
    assert cd.startswith('attachment; filename="kapsulka-dlya-vracha-') and "Для врача — Никита" in cd
    assert r.headers["cache-control"] == "no-store"


def test_xlsx_own_history_with_comments_and_disclaimer(client, fam):
    wb = xlsx(client, fam["h"], fam["f"])
    assert wb.sheetnames == ["Сводка", "Журнал приёма"]
    summary, log = wb["Сводка"], wb["Журнал приёма"]
    assert summary["A2"].value == DISCLAIMER and log["A2"].value == DISCLAIMER
    assert "Никита" in summary["A1"].value
    [(name, dosage, _, ingredient, times, total, unit, first, last)] = rows(summary)
    assert (name, dosage, ingredient, times, total, unit) == ("Нурофен", "200 мг", "ибупрофен", 1, 1, "таб")
    assert isinstance(first, datetime)
    [(when, name, _, amount, unit, comment)] = rows(log)
    assert (name, amount, comment) == ("Нурофен", 1, "болела голова")


def test_other_members_history_hides_their_comments(client, fam):
    """Комментарий — личная заметка: владелец выгружает историю мамы, но её комментариев не видит."""
    wb = xlsx(client, fam["h"], fam["f"], member=fam["mom"]["id"])
    log = wb["Журнал приёма"]
    assert log["A3"].value == NO_COMMENTS_NOTE
    assert rows(log) == [(rows(log)[0][0], "Нурофен", "200 мг", 0.5, "таб")]
    # Мама свою выписку получает с комментарием
    [row] = rows(xlsx(client, fam["h_mom"], fam["f"])["Журнал приёма"])
    assert row[-1] == "мамин секрет"


def test_period_and_timezone(client, fam, db):
    """Период считается по часовому поясу пользователя: 23:30 UTC — это уже следующий день в Москве."""
    for i in db.scalars(select(Intake)):
        i.taken_at = i.last_at = datetime(2026, 3, 9, 23, 30, tzinfo=timezone.utc)
    db.commit()
    f, h = fam["f"], fam["h"]
    assert rows(xlsx(client, h, f, **{"from": "2026-03-10", "to": "2026-03-10"})["Журнал приёма"])[0][0] == datetime(2026, 3, 10, 2, 30)
    assert rows(xlsx(client, h, f, **{"from": "2026-03-09", "to": "2026-03-09"})["Журнал приёма"]) == []
    london_day = xlsx(client, h, f, **{"from": "2026-03-09", "to": "2026-03-09", "tz": "Europe/London"})
    assert rows(london_day["Журнал приёма"])[0][0] == datetime(2026, 3, 9, 23, 30)
    # По умолчанию — последние 30 дней: старая запись не попадает
    assert rows(xlsx(client, h, f)["Журнал приёма"]) == []


def test_cabinet_sheet_is_optional(client, fam):
    wb = xlsx(client, fam["h"], fam["f"], cabinet="true")
    [row] = rows(wb["Аптечка"])
    assert row[:6] == ("Нурофен", "200 мг", None, "ибупрофен", 18.5, "таб")
    r = client.get(f"/api/families/{fam['f']}/report.pdf?cabinet=true", headers=fam["h"])
    assert r.status_code == 200


def test_deleted_medicine_stays_in_report(client, fam):
    client.delete(f"/api/families/{fam['f']}/medicines/{fam['med']['id']}", headers=fam["h"])
    [row] = rows(xlsx(client, fam["h"], fam["f"])["Журнал приёма"])
    assert row[1] == "Нурофен" and row[2] is None


@pytest.mark.parametrize("params, code", [
    ({"from": "2026-05-01", "to": "2026-04-01"}, 422),
    ({"from": "2020-01-01", "to": "2026-01-01"}, 422),
    ({"tz": "Mars/Olympus"}, 422),
    ({"tz": "../../etc/passwd"}, 422),
    ({"member": 99999}, 404),
])
def test_bad_params(client, fam, params, code):
    assert client.get(f"/api/families/{fam['f']}/report.pdf", headers=fam["h"], params=params).status_code == code


def test_unknown_format_404(client, fam):
    assert client.get(f"/api/families/{fam['f']}/report.docx", headers=fam["h"]).status_code == 404


def test_stranger_is_not_a_member(client, fam):
    h_x, x = register(client, "x@example.com", "Чужой")
    # Чужой не может выгрузить историю участника чужой семьи и через свою семью
    assert client.get(f"/api/families/{fid(x)}/report.pdf", headers=h_x, params={"member": fam["u"]["id"]}).status_code == 404


def test_plus_only_when_billing_on(client, fam):
    f, h = fam["f"], fam["h"]
    assert client.put("/api/admin/billing", headers=h, json={"enabled": True}).status_code == 200
    for ext in ("pdf", "xlsx"):
        r = client.get(f"/api/families/{f}/report.{ext}", headers=h)
        assert r.status_code == 402 and r.headers["x-plus-feature"] == "export_pdf"
    # Простой текстовый список остаётся бесплатным
    assert client.get(f"/api/families/{f}/export.txt", headers=h).status_code == 200
    until = (datetime.now(timezone.utc) + timedelta(days=30)).isoformat()
    assert client.put(f"/api/admin/families/{f}/plan", headers=h, json={"plan": "plus", "plus_until": until}).status_code == 200
    assert client.get(f"/api/families/{f}/report.pdf", headers=h).status_code == 200
